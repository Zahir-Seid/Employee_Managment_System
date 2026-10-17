"""Generic tabular export (Excel / PDF) for the Reports Module (spec
§32). Deliberately separate from services/excel_generator.py and
services/pdf_generator.py, which build the older, differently-shaped
"Generate Reports" (Stage 6) document — a multi-section, cover-paged
document for a single event/period/thematic scope, not a plain data
table. Every report in the Reports Module already resolves to one
filtered table of rows on screen (Event Activity, Communications, the
Management Report's Event Summary) — export should reproduce exactly
that table, not be bent to fit the older document's layout, and this
module never calls into that older generator's code. Only the header
fill/font colors are shared, for visual consistency across every
EMS-generated file.
"""

import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

VALID_EXPORT_FORMATS = {"xlsx", "pdf"}


def _stringify(value):
    return "" if value is None else str(value)


def rows_to_xlsx(sheet_title, headers, rows):
    """`headers`: [(label, key, width)]. `rows`: list of dicts — each
    column pulled via `row.get(key)`, so callers that need a derived
    value (e.g. Communications' "Responsible" = completed_by or event
    owner) precompute it into the row dict before calling this.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Report")[:31]

    for col_idx, (label, _key, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (_label, key, _width) in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_stringify(row.get(key)))

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def rows_to_pdf(title, headers, rows):
    """Same `headers`/`rows` shape as rows_to_xlsx. Landscape letter —
    reporting tables tend to be wide (many columns); switching
    orientation per report would leave already-tight columns tighter
    still for no benefit.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()

    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    table_data = [[label for label, _key, _width in headers]]
    for row in rows:
        table_data.append([_stringify(row.get(key)) for _label, key, _width in headers])

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    elements.append(table)

    if not rows:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph("No records match the selected filters.", styles["Normal"]))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def build_export_response(export_format, filename_base, title, headers, rows):
    """Returns a Django HttpResponse for the given format, or None if
    `export_format` isn't one this module handles (caller should
    validate against VALID_EXPORT_FORMATS before calling this).
    """
    if export_format == "xlsx":
        buffer = rows_to_xlsx(title, headers, rows)
        response = HttpResponse(
            buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename_base}.xlsx"'
        return response
    if export_format == "pdf":
        buffer = rows_to_pdf(title, headers, rows)
        response = HttpResponse(buffer.read(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename_base}.pdf"'
        return response
    return None
